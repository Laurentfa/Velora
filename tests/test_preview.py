import io
import zipfile

import pytest
from mailia import preview


def make_xlsx():
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Ventes"; ws.append(["Produit", "Qté"]); ws.append(["PPC", 3])
    b = io.BytesIO(); wb.save(b); return b.getvalue()


def make_pptx():
    import pptx
    p = pptx.Presentation(); s = p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text = "Titre"; s.placeholders[1].text = "Point"
    b = io.BytesIO(); p.save(b); return b.getvalue()


def test_excel_and_powerpoint_previews():
    r = preview.render("a.xlsx", make_xlsx()); assert r["type"] == "html" and "Ventes" in r["html"] and "PPC" in r["html"]
    r = preview.render("a.pptx", make_pptx()); assert "Titre" in r["html"] and "Diapositive 1" in r["html"]
    assert "PPC" in preview.extract_text("a.xlsx", make_xlsx())


def test_zip_listing_does_not_extract():
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w") as z:
        z.writestr("dossier/a.txt", "hello" * 100); z.writestr("../evil.sh", "x")
    r = preview.render("a.zip", b.getvalue())
    assert "dossier/a.txt" in r["html"] and r["type"] == "html"


def test_decompression_bomb_refused(monkeypatch):
    monkeypatch.setattr(preview, "MAX_UNCOMPRESSED", 1_000_000)
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", b"\0" * 5_000_000)
    with pytest.raises(preview.PreviewError, match="volumineux"):
        preview.render("bombe.docx", b.getvalue())
    with pytest.raises(preview.PreviewError):
        preview.extract_text("bombe.docx", b.getvalue())


def test_bad_files_give_clear_messages():
    for name in ("x.docx", "x.xlsx", "x.pptx"):
        with pytest.raises(preview.PreviewError):
            preview.render(name, b"pas un vrai fichier")
    with pytest.raises(preview.PreviewError, match="trop volumineux"):
        preview.render("gros.zip", b"0" * (preview.MAX_BYTES + 1))
