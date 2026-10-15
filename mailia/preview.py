"""Aperçu des pièces jointes qui ne sont pas lisibles telles quelles : Word, Excel, PowerPoint, HEIC, SVG, ZIP.

Tout se fait en local, avec des bibliothèques Python embarquées (aucune application tierce). Le résultat est soit une
image, soit un document HTML que l'appli nettoie puis affiche dans un cadre isolé. Ce sont des aperçus de lecture :
ni mise en page exacte, ni macros, ni formules recalculées. Rien n'est exécuté, rien n'est extrait sur le disque.
"""
import base64
import html
import io
import zipfile
from datetime import date, datetime, time

MAX_BYTES = 20 * 1024 * 1024  # au-delà, l'aperçu est refusé (le fichier reste téléchargeable)
MAX_HTML = 3_000_000
MAX_SHEETS, MAX_ROWS, MAX_COLS = 15, 400, 30
MAX_SLIDES = 60
MAX_ZIP_ENTRIES = 1000

KINDS = {
    "docx": "word", "docm": "word",
    "xlsx": "excel", "xlsm": "excel",
    "pptx": "powerpoint", "pptm": "powerpoint",
    "heic": "heic", "heif": "heic",
    "svg": "svg",
    "zip": "zip",
}


class PreviewError(Exception):
    """Message destiné à l'utilisateur."""


MAX_UNCOMPRESSED = 300 * 1024 * 1024   # taille totale une fois décompressé
MAX_ENTRIES = 5000
MAX_RATIO = 250                        # un fichier qui gonfle de plus de 250 fois est une « bombe » de décompression


def guard_archive(data: bytes) -> None:
    """Word, Excel et PowerPoint sont des archives ZIP : on vérifie qu'elles ne cachent pas une bombe de décompression avant de les ouvrir."""
    try:
        infos = zipfile.ZipFile(io.BytesIO(data)).infolist()
    except zipfile.BadZipFile:
        raise PreviewError("Ce fichier n'est pas un document Office valide (ou c'est un ancien format .doc/.xls/.ppt).")
    total = sum(i.file_size for i in infos)
    if len(infos) > MAX_ENTRIES or total > MAX_UNCOMPRESSED or any(i.compress_size and i.file_size / i.compress_size > MAX_RATIO and i.file_size > 5_000_000 for i in infos):
        raise PreviewError("Ce fichier est anormalement volumineux une fois décompressé : il est refusé par prudence. Vous pouvez le télécharger.")

def kind_of(name: str) -> str | None:
    ext = (name or "").rsplit(".", 1)[-1].lower() if "." in (name or "") else ""
    return KINDS.get(ext)


def _e(v) -> str:
    return html.escape(str(v), quote=True)


def _human(n: int) -> str:
    for unit in ("o", "Ko", "Mo", "Go"):
        if n < 1024 or unit == "Go":
            return f"{n:.0f} {unit}" if unit == "o" else f"{n:.1f} {unit}".replace(".", ",")
        n /= 1024
    return f"{n} o"


# ---------- images ----------
def _heic(data: bytes) -> dict:
    try:
        import pillow_heif
        from PIL import Image
    except ImportError:
        raise PreviewError("L'aperçu des photos HEIC n'est pas disponible dans cette version.")
    pillow_heif.register_heif_opener()
    try:
        im = Image.open(io.BytesIO(data))
        im.load()
    except Exception:
        raise PreviewError("Cette photo HEIC n'a pas pu être lue.")
    im = im.convert("RGB")
    im.thumbnail((2400, 2400))
    out = io.BytesIO()
    im.save(out, "JPEG", quality=88)
    return {"type": "image", "data": out.getvalue(), "ext": "jpg"}


def _svg(data: bytes) -> dict:
    # affiché dans une balise <img> : le navigateur n'y exécute aucun script ni ne charge rien d'externe
    return {"type": "image", "data": data, "ext": "svg"}


# ---------- Word ----------
def _word(data: bytes) -> dict:
    try:
        import mammoth
        result = mammoth.convert_to_html(io.BytesIO(data))
    except Exception:
        raise PreviewError("Ce document Word n'a pas pu être lu (fichier abîmé, protégé par mot de passe ou ancien format .doc).")
    body = result.value
    if not body.strip():
        raise PreviewError("Ce document Word ne contient aucun texte à afficher.")
    return {"type": "html", "html": _page(body)}


# ---------- Excel ----------
def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Oui" if v else "Non"
    if isinstance(v, datetime):
        return f"{v:%d/%m/%Y}" if (v.hour, v.minute, v.second) == (0, 0, 0) else f"{v:%d/%m/%Y %H:%M}"
    if isinstance(v, date):
        return f"{v:%d/%m/%Y}"
    if isinstance(v, time):
        return f"{v:%H:%M}"
    if isinstance(v, float):
        return f"{v:.10g}".replace(".", ",")
    return str(v)


def _excel(data: bytes) -> dict:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise PreviewError("Ce classeur Excel n'a pas pu être lu (fichier abîmé, protégé par mot de passe ou ancien format .xls).")
    parts: list[str] = []
    try:
        names = wb.sheetnames
        for name in names[:MAX_SHEETS]:
            ws = wb[name]
            rows: list[list[str]] = []
            truncated = False
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_ROWS:
                    truncated = True
                    break
                rows.append([_cell(v) for v in row[:MAX_COLS]])
            while rows and not any(rows[-1]):
                rows.pop()
            width = max((len(r) for r in rows), default=0)
            while width and not any(len(r) >= width and r[width - 1] for r in rows):
                width -= 1  # colonnes vides à droite
            parts.append(f"<h3>{_e(name)}</h3>")
            if not rows or not width:
                parts.append('<p class="muted">Feuille vide.</p>')
                continue
            parts.append("<table>")
            for r in rows:
                cells = (r + [""] * width)[:width]
                parts.append("<tr>" + "".join(f'<td class="w">{_e(c)}</td>' if len(c) > 40 else f"<td>{_e(c)}</td>" for c in cells) + "</tr>")
            parts.append("</table>")
            if truncated:
                parts.append(f'<p class="muted">… aperçu limité aux {MAX_ROWS} premières lignes.</p>')
        if len(names) > MAX_SHEETS:
            parts.append(f'<p class="muted">… {len(names) - MAX_SHEETS} autre(s) feuille(s) non affichée(s).</p>')
    finally:
        wb.close()
    return {"type": "html", "html": _page("".join(parts), extra=_EXCEL_CSS)}


# ---------- PowerPoint ----------
def _shape_html(shape, out: list[str], depth: int = 0) -> None:
    kind = getattr(shape, "shape_type", None)
    if depth < 4 and hasattr(shape, "shapes"):  # groupe de formes
        for s in sorted(shape.shapes, key=lambda s: (s.top or 0, s.left or 0)):
            _shape_html(s, out, depth + 1)
        return
    if getattr(shape, "has_table", False) and shape.has_table:
        out.append("<table>")
        for row in shape.table.rows:
            out.append("<tr>" + "".join(f"<td>{_e(c.text)}</td>" for c in row.cells) + "</tr>")
        out.append("</table>")
        return
    if kind == 13:  # image
        try:
            img = shape.image
            if img.content_type in ("image/png", "image/jpeg", "image/gif") and len(img.blob) <= 2_500_000:
                out.append(f'<p><img src="data:{img.content_type};base64,{base64.b64encode(img.blob).decode()}"></p>')
        except Exception:
            pass
        return
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        paras = [p for p in shape.text_frame.paragraphs if "".join(r.text for r in p.runs).strip()]
        for p in paras:
            text = "".join(r.text for r in p.runs).strip()
            pad = f' style="margin-left:{min(p.level, 5) * 18}px"' if p.level else ""
            out.append(f"<p{pad}>{_e(text)}</p>")


def _powerpoint(data: bytes) -> dict:
    try:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
    except Exception:
        raise PreviewError("Cette présentation PowerPoint n'a pas pu être lue (fichier abîmé, protégé par mot de passe ou ancien format .ppt).")
    parts: list[str] = []
    slides = list(prs.slides)
    for n, slide in enumerate(slides[:MAX_SLIDES], 1):
        title = ""
        try:
            if slide.shapes.title is not None:
                title = slide.shapes.title.text_frame.text.strip()
        except Exception:
            pass
        body: list[str] = []
        for shape in sorted(slide.shapes, key=lambda s: (s.top or 0, s.left or 0)):
            if title and getattr(shape, "has_text_frame", False) and shape.has_text_frame and shape.text_frame.text.strip() == title:
                continue  # déjà affiché en titre
            _shape_html(shape, body)
        parts.append(f'<div class="slide"><div class="num">Diapositive {n} sur {len(slides)}</div>'
                     + (f"<h3>{_e(title)}</h3>" if title else "") + "".join(body))
        try:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        except Exception:
            notes = ""
        if notes:
            parts.append(f'<div class="notes"><b>Notes :</b> {_e(notes)}</div>')
        parts.append("</div>")
    if len(slides) > MAX_SLIDES:
        parts.append(f'<p class="muted">… {len(slides) - MAX_SLIDES} autre(s) diapositive(s) non affichée(s).</p>')
    if not parts:
        raise PreviewError("Cette présentation ne contient aucune diapositive.")
    return {"type": "html", "html": _page("".join(parts), extra=_PPT_CSS)}


# ---------- ZIP ----------
def _zip(data: bytes) -> dict:
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
        infos = z.infolist()
    except Exception:
        raise PreviewError("Cette archive ZIP n'a pas pu être lue.")
    files = [i for i in infos if not i.is_dir()]
    total = sum(i.file_size for i in files)
    rows = []
    for i in files[:MAX_ZIP_ENTRIES]:
        lock = " 🔒" if i.flag_bits & 0x1 else ""
        rows.append(f"<tr><td>{_e(i.filename)}{lock}</td><td class=\"r\">{_human(i.file_size)}</td></tr>")
    head = (f"<p><b>{len(files)}</b> fichier(s), <b>{_human(total)}</b> une fois décompressé"
            f" ({_human(len(data))} compressé).</p>")
    if any(i.flag_bits & 0x1 for i in files):
        head += '<p class="muted">🔒 : fichier protégé par mot de passe.</p>'
    more = f'<p class="muted">… {len(files) - MAX_ZIP_ENTRIES} autre(s) fichier(s) non listé(s).</p>' if len(files) > MAX_ZIP_ENTRIES else ""
    return {"type": "html", "html": _page(head + "<table>" + "".join(rows) + "</table>" + more, extra=".r{text-align:right;white-space:nowrap}")}


# ---------- mise en page commune ----------
_EXCEL_CSS = ("table{border-collapse:collapse;font-size:13px;margin:6px 0 14px}"
              "td{border:1px solid #d6e0ee;padding:3px 8px;vertical-align:top;white-space:nowrap;max-width:260px;overflow:hidden;text-overflow:ellipsis}"  # une cellule reste sur une ligne : jamais de lettres empilées
              "td.w{white-space:normal;min-width:260px;max-width:520px;overflow:visible}"  # les longs textes (commentaires) passent à la ligne dans une colonne raisonnable
              "tr:first-child td{background:#eef3fa;font-weight:600;position:sticky;top:0;white-space:nowrap}")
_PPT_CSS = (".slide{border:1px solid #d6e0ee;border-radius:12px;padding:14px 18px;margin:0 0 14px;background:#fbfcfe}"
            ".num{font-size:11px;color:#66768a;text-transform:uppercase;letter-spacing:.06em;margin-bottom:6px}"
            ".notes{font-size:12.5px;color:#5a6b80;border-top:1px dashed #d6e0ee;margin-top:8px;padding-top:6px}"
            "td{border:1px solid #d6e0ee;padding:3px 8px}table{border-collapse:collapse}")


def _page(body: str, extra: str = "") -> str:
    css = ("h3{margin:14px 0 6px}.muted{color:#66768a;font-size:12.5px}table{border-collapse:collapse;margin:6px 0}"
           "td,th{border:1px solid #d6e0ee;padding:3px 8px;vertical-align:top}" + extra)
    return f"<style>{css}</style>{body}"


def render(name: str, data: bytes) -> dict:
    """{"type": "image", "data", "ext"} ou {"type": "html", "html"} ; lève PreviewError (message lisible) sinon."""
    if len(data) > MAX_BYTES:
        raise PreviewError(f"Ce fichier est trop volumineux pour l'aperçu ({_human(len(data))}). Téléchargez-le pour l'ouvrir.")
    fn = {"word": _word, "excel": _excel, "powerpoint": _powerpoint, "heic": _heic, "svg": _svg, "zip": _zip}.get(kind_of(name) or "")
    if not fn:
        raise PreviewError("Pas d'aperçu pour ce type de fichier.")
    if kind_of(name) in ("word", "excel", "powerpoint"):
        guard_archive(data)
    return fn(data)


# ---------- texte brut, pour le résumé IA à la demande ----------
MAX_TEXT_CHARS = 60_000  # environ 20 000 jetons : plafonne le coût d'un résumé
_TEXT_EXT = {"txt", "csv", "tsv", "log", "md", "json", "xml", "ics", "eml"}


def extract_text(name: str, data: bytes) -> str:
    """Texte lisible d'une pièce jointe (Word, Excel, PowerPoint, PDF, texte). Chaîne vide si rien d'exploitable ; PreviewError si illisible."""
    if len(data) > MAX_BYTES:
        raise PreviewError(f"Ce fichier est trop volumineux pour être résumé ({_human(len(data))}).")
    ext = (name or "").rsplit(".", 1)[-1].lower() if "." in (name or "") else ""
    kind = KINDS.get(ext)
    if kind in ("word", "excel", "powerpoint"):
        guard_archive(data)
    try:
        if kind == "word":
            import mammoth
            text = mammoth.extract_raw_text(io.BytesIO(data)).value
        elif kind == "excel":
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            lines: list[str] = []
            for ws in wb.worksheets[:MAX_SHEETS]:
                lines.append(f"## Feuille : {ws.title}")
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i >= 2000:
                        break
                    cells = [_cell(v) for v in row[:MAX_COLS]]
                    if any(cells):
                        lines.append("\t".join(cells).rstrip())
            wb.close()
            text = "\n".join(lines)
        elif kind == "powerpoint":
            from pptx import Presentation
            lines = []
            for n, slide in enumerate(Presentation(io.BytesIO(data)).slides, 1):
                lines.append(f"## Diapositive {n}")
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False) and shape.has_text_frame and shape.text_frame.text.strip():
                        lines.append(shape.text_frame.text.strip())
                    elif getattr(shape, "has_table", False) and shape.has_table:
                        lines += ["\t".join(c.text for c in r.cells) for r in shape.table.rows]
                try:
                    if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                        lines.append("Notes : " + slide.notes_slide.notes_text_frame.text.strip())
                except Exception:
                    pass
            text = "\n".join(lines)
        elif ext == "pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise PreviewError("Ce PDF est protégé par un mot de passe.")
            chunks, size = [], 0
            for page in reader.pages[:80]:
                t = page.extract_text() or ""
                chunks.append(t)
                size += len(t)
                if size > MAX_TEXT_CHARS:
                    break
            text = "\n".join(chunks)
        elif ext in _TEXT_EXT:
            text = data[:MAX_TEXT_CHARS * 2].decode("utf-8", errors="replace")
        else:
            raise PreviewError("Je ne sais résumer que les documents Word, Excel, PowerPoint, PDF et les fichiers texte.")
    except PreviewError:
        raise
    except Exception:
        raise PreviewError("Le texte de ce fichier n'a pas pu être lu (fichier abîmé ou protégé).")
    text = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    return text[:MAX_TEXT_CHARS]


# ---------- images d'une pièce jointe, pour la lecture d'un scan par l'IA ----------
MAX_VISION_PAGES = 6


def scan_images(name: str, data: bytes) -> tuple[list[bytes], bool]:
    """(images JPEG réduites, tronqué ?) d'un PDF numérisé ou d'une photo, pour la lecture par l'IA."""
    if len(data) > MAX_BYTES:
        raise PreviewError(f"Ce fichier est trop volumineux pour être lu par l'IA ({_human(len(data))}).")
    from PIL import Image, ImageOps
    ext = (name or "").rsplit(".", 1)[-1].lower() if "." in (name or "") else ""
    raws: list = []
    truncated = False
    try:
        if ext == "pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise PreviewError("Ce PDF est protégé par un mot de passe.")
            truncated = len(reader.pages) > MAX_VISION_PAGES
            for page in reader.pages[:MAX_VISION_PAGES]:
                imgs = list(page.images)
                if imgs:
                    raws.append(max(imgs, key=lambda i: len(i.data)).image)  # l'image la plus lourde de la page : le scan
        elif ext in ("heic", "heif"):
            import pillow_heif
            pillow_heif.register_heif_opener()
            raws.append(Image.open(io.BytesIO(data)))
        else:
            raws.append(Image.open(io.BytesIO(data)))
        out = []
        for im in raws:
            im = ImageOps.exif_transpose(im).convert("RGB")
            im.thumbnail((1600, 1600))
            buf = io.BytesIO()
            im.save(buf, "JPEG", quality=78)
            out.append(buf.getvalue())
    except PreviewError:
        raise
    except Exception:
        raise PreviewError("Les images de ce fichier n'ont pas pu être lues.")
    return out, truncated


def csv_table(data: bytes, max_rows: int = 500, max_cols: int = 60) -> tuple[str, int, int] | None:
    """Fichier CSV/TSV -> tableau HTML (échappé). Retourne (html, lignes au total, lignes affichées), ou None si ce n'est pas un tableau lisible."""
    import csv
    for enc in ("utf-8-sig", "cp1252"):
        try:
            text = data[:3_000_000].decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        return None
    sample = text[:20000]
    try:
        delim = csv.Sniffer().sniff(sample, delimiters=";,\t|").delimiter
    except csv.Error:
        delim = max(";,\t|", key=sample.count)
    rows = []
    total = 0
    for r in csv.reader(io.StringIO(text), delimiter=delim):
        if not any(c.strip() for c in r):
            continue
        total += 1
        if len(rows) < max_rows:
            rows.append([c[1:].strip('"') if c.startswith('="') else c for c in r[:max_cols]])  # Excel écrit les nombres « ="0123" »
    if not rows or max(len(r) for r in rows) < 2:
        return None
    width = max(len(r) for r in rows)
    cell = lambda c: html.escape(c if len(c) <= 400 else c[:400] + "…")
    head = "".join(f"<th>{cell(c)}</th>" for c in rows[0])
    body = "".join("<tr>" + "".join(f"<td>{cell(c)}</td>" for c in r + [""] * (width - len(r))) + "</tr>" for r in rows[1:])
    return f"<table class='csvt'><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>", max(total - 1, 0), len(rows) - 1
